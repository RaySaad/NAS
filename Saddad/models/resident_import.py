import base64
import io
import json
import zipfile
from datetime import date, datetime
from collections import Counter
import openpyxl
from hijridate import Gregorian
from odoo import api, fields, models, _
from odoo.exceptions import AccessError, UserError, ValidationError
from .resident import Resident, identity


HEADERS = ['رقم الاقامة', 'الاسم', 'الجنس', 'الجنسية', 'المهنة', 'رقم الجواز',
           'تاريخ انتهاء الجواز', 'تاريخ اصدار الاقامة', 'تاريخ انتهاء الاقامة',
           'تاريخ الميلاد', 'خارج المملكه', 'تاريخ انتهاء الاقامة بالهجري', 'رقم صاحب العمل']
KEYS = ['iqama', 'reported_name', 'sex', 'nationality', 'occupation', 'passport',
        'passport_expiry', 'iqama_issue', 'reported_expiry', 'birth_date',
        'outside_kingdom', 'reported_hijri', 'employer']
DATE_KEYS = {'passport_expiry', 'iqama_issue', 'reported_expiry', 'birth_date'}


def date_value(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return date.fromisoformat(str(value).strip()).isoformat()


class ResidentImport(models.Model):
    _name = 'saddad.resident.import'
    _description = 'Muqeem Resident Import'
    _inherit = ['mail.thread']
    _rec_name = 'filename'
    company_id = fields.Many2one('res.company', required=True, default=lambda s: s.env.company)
    filename = fields.Char(required=True)
    file = fields.Binary(required=True, attachment=True)
    snapshot_date = fields.Date(required=True, help='Date the report was exported from Muqeem, confirmed by HR.')
    state = fields.Selection([('draft', 'Draft'), ('preview', 'Preview'), ('applied', 'Applied')], default='draft', readonly=True, tracking=True)
    line_ids = fields.One2many('saddad.resident.import.line', 'batch_id', readonly=True)
    summary = fields.Text(readonly=True)
    applied_at = fields.Datetime(readonly=True)
    applied_by = fields.Many2one('res.users', readonly=True)

    @api.model_create_multi
    def create(self, vals_list):
        if any({'state', 'line_ids', 'summary', 'applied_at', 'applied_by'}.intersection(v) for v in vals_list):
            raise AccessError(_('Import outcomes are managed by import actions.'))
        return super().create([dict(v, state='draft', summary=False, applied_at=False, applied_by=False, line_ids=False) for v in vals_list])

    def write(self, vals):
        if {'state', 'line_ids', 'summary', 'applied_at', 'applied_by'}.intersection(vals) or any(r.state != 'draft' for r in self):
            raise AccessError(_('Create a new import to change a reviewed report.'))
        return super().write(vals)

    def unlink(self):
        raise UserError(_('Import history cannot be deleted.'))

    def _lock(self):
        self.ensure_one()
        self.check_access('write')
        self.env.cr.execute('SELECT id FROM saddad_resident_import WHERE id = %s FOR UPDATE', [self.id])
        self.invalidate_recordset()

    def action_preview(self):
        self._lock()
        if self.state != 'draft':
            raise UserError(_('This import has already been previewed.'))
        if self.snapshot_date > fields.Date.context_today(self):
            raise UserError(_('The snapshot date cannot be in the future.'))
        try:
            data = base64.b64decode(self.file, validate=True)
            if len(data) > 15 * 1024 * 1024:
                raise ValueError()
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                if sum(f.file_size for f in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError()
            book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            sheet = book.active
            if sheet.max_row > 25001 or sheet.max_column != 13:
                raise ValueError()
            rows = sheet.iter_rows(values_only=True)
            headers = [str(v or '').strip() for v in next(rows)]
            if set(headers) != set(HEADERS):
                raise ValueError()
            order = [headers.index(h) for h in HEADERS]
            records = [(n, [r[i] for i in order]) for n, r in enumerate(rows, 2) if any(v is not None for v in r)]
            book.close()
        except Exception as exc:
            raise UserError(_('Upload the active-residents XLSX report with its 13 original columns, no more than 25,000 rows and 15 MB.')) from exc
        identities = Counter(identity(r[0]) for _, r in records)
        employees = self.env['hr.employee'].with_context(active_test=False).search([('company_id', '=', self.company_id.id)])
        matches = {}
        for employee in employees:
            matches.setdefault(identity(employee.identification_id), []).append(employee)
        establishments = {r.number: r for r in self.env['saddad.establishment'].search([('company_id', '=', self.company_id.id)])}
        profiles = {r.employee_id.id: r for r in self.env['saddad.resident'].search([('company_id', '=', self.company_id.id)])}
        labels = self.env['saddad.resident'].fields_get(attributes=['string'])
        prepared = []
        for number, row in records:
            values, error = {}, ''
            employee = False
            profile = False
            try:
                iqama, employer = identity(row[0]), identity(row[12])
                if len(iqama) != 10 or not iqama.isdigit() or identities[iqama] > 1:
                    raise ValueError(_('Missing, invalid or duplicated Iqama number.'))
                if employer not in establishments:
                    raise ValueError(_('Employer number is not mapped to this company.'))
                candidates = matches.get(iqama, [])
                if len(candidates) != 1 or not candidates[0].active:
                    raise ValueError(_('No unique active employee matches this Iqama in this company.'))
                employee = candidates[0]
                profile = profiles.get(employee.id)
                if profile and profile.snapshot_date > self.snapshot_date:
                    raise ValueError(_('This report is older than the stored resident snapshot.'))
                for key, raw in zip(KEYS, row):
                    if key == 'employer' or raw is None or str(raw).strip() == '':
                        continue
                    values[key] = date_value(raw) if key in DATE_KEYS else str(raw).strip()
                values['iqama'] = iqama
                if not values.get('reported_expiry'):
                    raise ValueError(_('The report must contain an Iqama expiry date.'))
                expiry = fields.Date.to_date(values['reported_expiry'])
                hijri = Gregorian(expiry.year, expiry.month, expiry.day).to_hijri().isoformat()
                if values.get('reported_hijri') and identity(values['reported_hijri']) != hijri:
                    raise ValueError(_('Gregorian and Hijri expiry dates do not agree. Review the source report.'))
                values.update(company_id=self.company_id.id, employee_id=employee.id,
                              establishment_id=establishments[employer].id, snapshot_date=str(self.snapshot_date))
            except (ValueError, OverflowError) as exc:
                error = str(exc)
            changes = []
            if not error:
                for key, value in values.items():
                    current = profile[key] if profile else False
                    if hasattr(current, 'id'):
                        current = current.id
                    if str(current or '') != str(value):
                        descriptor = self.env['saddad.resident']._fields[key]
                        old_display, new_display = current or '', value
                        if descriptor.type == 'many2one':
                            old_display = profile[key].display_name if profile and profile[key] else ''
                            new_display = self.env[descriptor.comodel_name].browse(value).display_name
                        changes.append('%s: %s → %s' % (labels[key]['string'], old_display, new_display))
            prepared.append({'batch_id': self.id, 'row_number': number, 'iqama': identity(row[0]),
                'employee_id': employee.id if employee else False, 'status': 'error' if error else 'ready',
                'detail': error or '\n'.join(changes) or _('No changes'),
                'payload': json.dumps({'values': values, 'profile_id': profile.id if profile else False,
                    'version': str(profile.write_date) if profile else False}),
                'proposed_expiry': values.get('reported_expiry') if not error else False})
        super(ResidentImportLine, self.env['saddad.resident.import.line']).create(prepared)
        errors = sum(v['status'] == 'error' for v in prepared)
        super().write({'state': 'preview', 'summary': _('Rows: %(rows)s. Ready: %(ready)s. Blocked: %(errors)s. Applying imports ready rows only; blocked rows remain visible. Employee fields are not changed.', rows=len(prepared), ready=len(prepared)-errors, errors=errors)})

    def action_apply(self):
        self._lock()
        if self.state != 'preview':
            raise UserError(_('Preview the import before applying it.'))
        ready = self.line_ids.filtered(lambda l: l.status == 'ready')
        if not ready:
            raise UserError(_('There are no ready rows to import.'))
        # A company lock serializes imports even where a resident does not exist yet.
        self.env.cr.execute('SELECT id FROM res_company WHERE id = %s FOR UPDATE', [self.company_id.id])
        for line in ready:
            payload = json.loads(line.payload)
            values = payload['values']
            employee = line.employee_id
            employee.check_access('read')
            if employee.company_id != self.company_id or not employee.active or identity(employee.identification_id) != line.iqama:
                raise UserError(_('Employee matching changed. Create a fresh import preview.'))
            profile = self.env['saddad.resident'].search([('employee_id', '=', employee.id)], limit=1)
            if profile:
                self.env.cr.execute('SELECT id FROM saddad_resident WHERE id = %s FOR UPDATE', [profile.id])
                profile.invalidate_recordset()
            if (profile.id or False) != payload['profile_id'] or (str(profile.write_date) if profile else False) != payload['version']:
                raise UserError(_('Resident data changed since preview. Create a fresh import preview.'))
            establishment = self.env['saddad.establishment'].browse(values['establishment_id'])
            if establishment.company_id != self.company_id:
                raise UserError(_('Establishment mapping changed. Create a fresh import preview.'))
            values['source_batch_id'] = self.id
            if profile:
                super(Resident, profile).write(values)
            else:
                profile = super(Resident, self.env['saddad.resident']).create(values)
            profile.message_post(body=_('Resident snapshot imported from %s.', self.filename))
            super(ResidentImportLine, line).write({'status': 'applied'})
        super().write({'state': 'applied', 'applied_at': fields.Datetime.now(), 'applied_by': self.env.uid})


class ResidentImportLine(models.Model):
    _name = 'saddad.resident.import.line'
    _description = 'Muqeem Import Preview Row'
    _order = 'row_number'
    batch_id = fields.Many2one('saddad.resident.import', required=True, ondelete='restrict')
    company_id = fields.Many2one(related='batch_id.company_id', store=True)
    row_number = fields.Integer()
    iqama = fields.Char()
    employee_id = fields.Many2one('hr.employee')
    proposed_expiry = fields.Date()
    status = fields.Selection([('ready', 'Ready'), ('error', 'Blocked'), ('applied', 'Applied')])
    detail = fields.Text()
    payload = fields.Text()

    @api.model_create_multi
    def create(self, vals_list):
        raise AccessError(_('Preview rows are generated by the import.'))

    def write(self, vals):
        raise AccessError(_('Preview rows cannot be edited. Create a corrected import.'))

    def unlink(self):
        raise AccessError(_('Import history cannot be deleted.'))
