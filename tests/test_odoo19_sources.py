"""Source and isolated-method regressions; these are NOT Odoo runtime tests."""
import ast
import base64
import copy
import datetime
import importlib.util
import json
import io
from pathlib import Path
from types import SimpleNamespace as NS
import unittest
import subprocess
import sys

from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
BACKUP = ROOT.parent / 'NAS (copy)'
if not BACKUP.is_dir():
    BACKUP = ROOT.parent / 'NAS_old'


def load_methods(filename, name, methods, base=object, **namespace):
    tree = ast.parse((ROOT / filename).read_bytes())
    cls = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == name)
    cls = copy.deepcopy(cls)
    cls.bases = [ast.Name(id='Base', ctx=ast.Load())]
    cls.keywords = []
    cls.decorator_list = []
    cls.body = [n for n in cls.body if isinstance(n, ast.FunctionDef) and n.name in methods]
    for method in cls.body:
        method.decorator_list = []
    namespace['Base'] = base
    module = ast.fix_missing_locations(ast.Module(body=[cls], type_ignores=[]))
    exec(compile(module, filename, 'exec'), namespace)
    return namespace[name]


def field_definitions(path):
    result = set()
    for cls in ast.parse(path.read_bytes()).body:
        if not isinstance(cls, ast.ClassDef):
            continue
        for node in cls.body:
            if not isinstance(node, ast.Assign) or not isinstance(node.value, ast.Call):
                continue
            call = node.value.func
            if isinstance(call, ast.Attribute) and isinstance(call.value, ast.Name) and call.value.id == 'fields':
                result.update((cls.name, t.id) for t in node.targets if isinstance(t, ast.Name))
    return result


class SourceRegressionTests(unittest.TestCase):
    @unittest.skipUnless(BACKUP.is_dir(), 'Old contacts API source comparison requires backup.')
    def test_contacts_api_business_logic_matches_old(self):
        relative = 'contacts_ext/controllers/contact_api.py'
        original = ast.parse((BACKUP / relative).read_bytes())
        current = ast.parse((ROOT / relative).read_bytes())
        # These are the only approved migration differences in this module:
        # JSON route rename and flat/JSON-RPC payload compatibility helper.
        current.body = [n for n in current.body if not (
            isinstance(n, ast.ImportFrom) and n.module == 'api_utils')]
        class NormalizePayload(ast.NodeTransformer):
            def visit_Call(self, node):
                if isinstance(node.func, ast.Name) and node.func.id == 'parse_api_payload':
                    return ast.parse("json.loads(request.httprequest.data.decode('utf-8'))", mode='eval').body
                return self.generic_visit(node)
        current = NormalizePayload().visit(current)
        for node in ast.walk(original):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == 'route':
                for option in node.keywords:
                    if option.arg == 'type' and isinstance(option.value, ast.Constant) and option.value.value == 'json':
                        option.value.value = 'jsonrpc'
        self.assertEqual(ast.dump(current, include_attributes=False), ast.dump(original, include_attributes=False))

    @unittest.skipUnless(BACKUP.is_dir(), 'Old API source comparison requires backup.')
    def test_account_api_matches_old_except_required_route_type(self):
        relative = 'account_ext/controllers/jv_api.py'
        original = ast.parse((BACKUP / relative).read_bytes())
        current = ast.parse((ROOT / relative).read_bytes())
        # Odoo 19 renamed the JSON route type. Every other statement, route
        # option, payload parser and hardcoded journal must match production.
        for node in ast.walk(original):
            if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == 'route':
                for option in node.keywords:
                    if option.arg == 'type' and isinstance(option.value, ast.Constant) and option.value.value == 'json':
                        option.value.value = 'jsonrpc'
        self.assertEqual(ast.dump(current, include_attributes=False),
                         ast.dump(original, include_attributes=False))

    def test_employee_native_version_storage_and_custom_code_preserved(self):
        tree = ast.parse((ROOT / 'employees_ext/models/employee.py').read_text())
        cls = next(n for n in tree.body if isinstance(n, ast.ClassDef) and n.name == 'EmployeeExt')
        declarations = {n.targets[0].id: n.value for n in cls.body
                        if isinstance(n, ast.Assign) and isinstance(n.targets[0], ast.Name)
                        and isinstance(n.value, ast.Call)}
        for name in ('identification_id', 'marital'):
            options = {kw.arg: ast.literal_eval(kw.value) for kw in declarations[name].keywords}
            self.assertEqual(options['related'], 'version_id.' + name)
            self.assertIs(options['store'], False)
            self.assertIs(options['readonly'], False)
            self.assertIs(options['inherited'], True)
        self.assertEqual(declarations['employee_code'].func.attr, 'Char')
        options = {kw.arg: ast.literal_eval(kw.value) for kw in declarations['employee_code'].keywords}
        self.assertIs(options['required'], True)
        self.assertEqual(declarations['_employee_code_unique'].args[0].value, 'unique(employee_code)')

    def test_contact_layout_with_odoo19_inheritance_engine(self):
        odoo_root = Path('/home/ayyan-dev/odoo19/odoo')
        if not (odoo_root / 'odoo/addons/base/views/res_partner_views.xml').is_file():
            self.skipTest('Local Odoo 19 source is unavailable')
        result = subprocess.run([
            sys.executable, '-B', str(ROOT / 'tools/check_odoo19_contact_layout.py'),
            '--odoo-root', str(odoo_root),
        ], capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stdout + result.stderr)

    @unittest.skipUnless(BACKUP.is_dir(), 'Backup comparison requires NAS (copy).')
    def test_original_voucher_buttons_retained(self):
        for module in ('custom_journal_excel_export', 'journal_voucher_print_export', 'journal_entry_print'):
            relative = Path(module) / 'views/account_move_views.xml'
            original = etree.parse(str(BACKUP / relative)).xpath('//button')
            current = etree.parse(str(ROOT / relative)).xpath('//button')
            self.assertEqual([b.get('string') for b in current], [b.get('string') for b in original])
            self.assertEqual([b.get('type') for b in current], [b.get('type') for b in original])

    def test_all_python_and_xml_syntax(self):
        for path in ROOT.glob('*/**/*.py'):
            with self.subTest(path=path):
                ast.parse(path.read_bytes(), filename=str(path))
        for path in ROOT.glob('*/**/*.xml'):
            with self.subTest(path=path):
                etree.parse(str(path))

    @unittest.skipUnless(BACKUP.is_dir(), 'Backup comparison requires NAS (copy).')
    def test_backup_files_and_custom_fields_preserved(self):
        for path in BACKUP.rglob('*'):
            if not path.is_file() or any(p in ('.git', '__pycache__') for p in path.parts) or path.suffix == '.pyc':
                continue
            current = ROOT / path.relative_to(BACKUP)
            with self.subTest(path=path):
                self.assertTrue(current.is_file(), 'Original file is missing')
                if path.suffix == '.py':
                    self.assertFalse(field_definitions(path) - field_definitions(current), 'Original custom fields removed')

    @unittest.skipUnless(BACKUP.is_dir(), 'Backup comparison requires NAS (copy).')
    def test_original_api_routes_preserved(self):
        def routes(path):
            result = set()
            for node in ast.walk(ast.parse(path.read_bytes())):
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == 'route' and node.args:
                    route = ast.literal_eval(node.args[0])
                    result.update([route] if isinstance(route, str) else route)
            return result
        for path in BACKUP.glob('*/controllers/*.py'):
            current = ROOT / path.relative_to(BACKUP)
            self.assertFalse(routes(path) - routes(current), str(path))

    @unittest.skipUnless(BACKUP.is_dir(), 'Backup comparison requires NAS (copy).')
    def test_sql_constraints_preserved(self):
        for path in BACKUP.glob('*/models/*.py'):
            old = ast.parse(path.read_bytes())
            expected = []
            for node in ast.walk(old):
                if isinstance(node, ast.Assign) and any(isinstance(t, ast.Name) and t.id == '_sql_constraints' for t in node.targets):
                    expected.extend(ast.literal_eval(node.value))
            if not expected:
                continue
            actual = {}
            for node in ast.walk(ast.parse((ROOT / path.relative_to(BACKUP)).read_bytes())):
                if isinstance(node, ast.Assign) and isinstance(node.value, ast.Call) and isinstance(node.value.func, ast.Attribute) and node.value.func.attr == 'Constraint':
                    actual[node.targets[0].id.lstrip('_')] = tuple(ast.literal_eval(a) for a in node.value.args[:2])
            for name, sql, message in expected:
                self.assertEqual(actual.get(name), (sql, message), str(path))

    def test_no_deprecated_json_routes_or_user_group_fields(self):
        module_paths = [p.parent for p in ROOT.glob('*/__manifest__.py')]
        for path in (p for module in module_paths for p in module.rglob('*.py')):
            tree = ast.parse(path.read_bytes())
            for node in ast.walk(tree):
                if isinstance(node, ast.Constant):
                    self.assertNotEqual(node.value, 'groups_id', str(path))
                if isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute) and node.func.attr == 'route':
                    for keyword in node.keywords:
                        if keyword.arg == 'type':
                            self.assertNotEqual(ast.literal_eval(keyword.value), 'json', str(path))
        for path in ROOT.glob('*/**/*.xml'):
            for record in etree.parse(str(path)).xpath('//record[@model="res.groups"]'):
                self.assertFalse(record.xpath('./field[@name="users" or @name="category_id"]'), str(path))

    def test_legacy_mobile_column_preserved(self):
        self.assertIn(('ResPartner', 'mobile'), field_definitions(ROOT / 'contacts_ext/models/contact.py'))

    def test_independent_ou_permissions_have_independent_privileges(self):
        files = ['operating_unit/security/operating_unit_security.xml',
                 'operating_unit_access_all/security/operating_unit_security.xml',
                 'stock_operating_unit_access_all/security/stock_security.xml']
        privileges = []
        for filename in files:
            tree = etree.parse(str(ROOT / filename))
            privileges.append(tree.xpath('//record[@model="res.groups"]/field[@name="privilege_id"]')[0].get('ref'))
        self.assertEqual(len(set(privileges)), 3)


class IsolatedMethodTests(unittest.TestCase):
    def test_version_marital_choices_keep_core_and_custom_other(self):
        class Base:
            def _get_marital_status_selection(self):
                return [('single', 'Single'), ('married', 'Married')]
        cls = load_methods('employees_ext/models/employee.py', 'EmployeeVersionExt',
                           {'_get_marital_status_selection'}, Base, _=lambda text: text)
        self.assertEqual(cls()._get_marital_status_selection(),
                         [('single', 'Single'), ('married', 'Married'), ('other', 'Other')])

    def test_legacy_iap_settings_button_uses_odoo19_action(self):
        cls = load_methods('employees_ext/models/token_settings.py', 'ResConfigSettings',
                           {'iap_action_view_my_services'})
        settings = cls()
        called = []
        action = {'type': 'ir.actions.act_window', 'res_model': 'iap.account'}
        def resolve(xml_id):
            called.append(xml_id)
            return action
        settings.env = {'ir.actions.actions': NS(_for_xml_id=resolve)}
        self.assertIs(settings.iap_action_view_my_services(), action)
        self.assertEqual(called, ['iap.iap_account_action'])
        manifest = ast.literal_eval((ROOT / 'employees_ext/__manifest__.py').read_text())
        self.assertIn('iap', manifest['depends'])

    def test_petty_cash_split_tax_not_duplicated(self):
        class Base:
            def ensure_one(self):
                pass
            def write(self, values):
                self.written = values
        cls = load_methods('petty_cash_purchases/models/petty_cash_purchases.py', 'PettyCashPurchase', {'post'}, Base,
                           ValidationError=ValueError, _=lambda text: text,
                           fields=NS(Datetime=NS(now=lambda: datetime.datetime(2026, 10, 1))))
        purchase = cls()
        tax_rows = [{'id': 3, 'name': 'VAT', 'account_id': 4, 'tax_repartition_line_id': rid,
                     'amount': 5, 'tag_ids': []} for rid in (5, 6)]
        taxes = NS(ids=[3], compute_all=lambda **kwargs: {'taxes': tax_rows, 'base_tags': []})
        line = NS(tax_ids=taxes, price_unit=100, quantity=1, company_currency_id=NS(id=1),
                  name='Expense', analytic_distribution={'10,20': 100}, account_id=NS(id=2),
                  operating_unit_id=NS(id=8), vehicle_id=NS(id=False), vendor_id=NS(id=7),
                  employee_id=NS(id=9), customer_account=NS(id=10), sub_total=100,
                  customer_code='C1', contract_type='Type', employee_code='E1')
        captured = []
        def create(values):
            captured.append(values)
            return NS(id=42)
        purchase.env = {'account.move': NS(create=create)}
        purchase.expenses_line_ids = [line]
        purchase.account_id, purchase.company_id, purchase.employee_journal_id = NS(id=11), NS(id=1), NS(id=12)
        purchase.name, purchase.id = 'PCP1', 1
        purchase.post()
        values = [values for _, _, values in captured[0]['line_ids']]
        self.assertEqual(len(values), 4)
        self.assertEqual([v['debit'] for v in values if v.get('tax_line_id')], [5, 5])
        self.assertEqual(values[-1]['credit'], 110)
        self.assertEqual(sum(v.get('debit', 0)-v.get('credit', 0) for v in values), 0)

    def test_real_voucher_workbooks_formulas_and_analytic_labels(self):
        import openpyxl
        import xlsxwriter
        class Base:
            def ensure_one(self):
                pass
        class Lines(list):
            def mapped(self, field):
                return [getattr(line, field) for line in self]
        for module, method in [('custom_journal_excel_export', 'export_large_voucher_excel'),
                               ('journal_voucher_print_export', 'export_voucher_excel')]:
            with self.subTest(module=module):
                cls = load_methods(module + '/models/account_move.py', 'AccountMove', {method}, Base,
                                   xlsxwriter=xlsxwriter, BytesIO=io.BytesIO, base64=base64,
                                   UserError=ValueError, _=lambda text: text,
                                   _logger=NS(info=lambda *a: None, warning=lambda *a: None, error=lambda *a: None))
                move = cls()
                attachments = []
                def create(values):
                    attachments.append(values)
                    return NS(id=1)
                move.env = {'ir.attachment': NS(create=create),
                            'account.analytic.account': NS(browse=lambda ids: NS(mapped=lambda field: ['Plan A', 'Plan B']))}
                move.company_id, move.id, move.name = NS(name='Test Company'), 42, 'JV/42'
                move.date, move.journal_id, move.ref, move.state = datetime.date(2026, 10, 1), NS(name='General'), '', 'posted'
                move._fields = {'state': NS(selection=[('posted', 'Posted')])}
                def line(debit, credit):
                    return NS(account_id=NS(code='00123', name='Account'), partner_id=NS(name='Partner'),
                              employee_id=NS(name='Employee'), name='Test line', debit=debit, credit=credit,
                              analytic_distribution={'10,20': 100}, id=1)
                move.line_ids = Lines([line(100, 0), line(0, 100)])
                result = getattr(move, method)()
                self.assertEqual(result['type'], 'ir.actions.act_url')
                workbook = openpyxl.load_workbook(io.BytesIO(base64.b64decode(attachments[0]['datas'])))
                sheet = workbook.active
                total_row = next(row[4].row for row in sheet if row[4].value == 'TOTAL')
                self.assertEqual(sheet.cell(total_row, 6).value, f'=SUM(F{total_row-2}:F{total_row-1})')
                self.assertEqual(sheet.cell(total_row, 7).value, f'=SUM(G{total_row-2}:G{total_row-1})')
                self.assertEqual(sheet.cell(total_row+2, 5).value, f'=F{total_row}-G{total_row}')
                self.assertEqual(sheet.cell(total_row-1, 8).value, 'Plan A, Plan B')
                self.assertEqual(sheet.cell(total_row-1, 1).value, '00123')
                workbook.close()

    def test_statement_opening_and_period_totals(self):
        class Base:
            def ensure_one(self):
                pass
        cls = load_methods('account_ext/wizard/full_employee_statement_wizard.py', 'FullEmployeeStatementWizard',
                           {'_base_domain', 'action_generate_statement'}, Base, _=lambda text: text)
        wizard = cls()
        employee1, employee2, account = NS(id=1, employee_code='E1'), NS(id=2, employee_code='E2'), NS(id=7)
        captured = []
        def read_group(**kwargs):
            self.assertIn(('company_id', '=', 3), kwargs['domain'])
            if kwargs['aggregates'] == ['debit:sum', 'credit:sum']:
                return [(employee1, account, 10, 3)]
            return [(employee1, account, 5), (employee2, account, 20)]
        class Env(dict):
            context = {}
            def ref(self, name):
                return NS(id=99)
        wizard.env = Env({'account.move.line': NS(_read_group=read_group),
                          'hr.employee': NS(browse=lambda ids: [employee1, employee2]),
                          'full.employee.statement.line': NS(create=captured.extend)})
        wizard.company_id, wizard.account_ids = NS(id=3), NS(ids=[7])
        wizard.date_from, wizard.date_to = '2026-01-01', '2026-12-31'
        wizard.line_ids, wizard.id, wizard._name = NS(unlink=lambda: None), 5, 'full.employee.statement.wizard'
        result = wizard.action_generate_statement()
        self.assertEqual(result['res_id'], 5)
        self.assertEqual([(r['employee_code'], r['balance']) for r in captured], [('E1', 12), ('E2', 20)])

    def test_excel_lookup_cache_leading_zeros_and_employee_precedence(self):
        cls = load_methods('account_ext/models/excel_jv_import_wizard.py', 'ExcelJVImportWizard',
                           {'_process_line_item', '_map_header_to_field', '_normalize_field_name'},
                           ValidationError=ValueError, _=lambda text: text, json=json, _logger=NS(debug=lambda *a: None))
        importer = cls()
        partner = NS(id=7, customer_code='C1')
        employee = NS(id=9, customer_account=NS(id=6, partner_id=partner), operating_unit_id=NS(id=8))
        employee.with_context = lambda **kwargs: employee
        searches = []
        def search_account(domain, **kwargs):
            searches.append(domain)
            return NS(id=1)
        importer.env = {'account.account': NS(search=search_account),
                        'hr.employee': NS(search=lambda *a, **k: employee)}
        caches = {key: {} for key in ('account', 'operating_unit', 'customer_account', 'vehicle', 'employee')}
        fields = dict.fromkeys(['account_id', 'employee_code', 'partner_id', 'customer_code', 'operating_unit_id'])
        row = {'Account Code': '00123', 'Employee Code': 'E1', 'partner_id': 100, 'customer_code': 'WRONG'}
        for _ in range(2):
            result = importer._process_line_item(row, fields, 1, caches)
            self.assertEqual(result['partner_id'], 7)
            self.assertEqual(result['customer_code'], 'C1')
            self.assertEqual(result['operating_unit_id'], 8)
        self.assertEqual(searches, [[('code', '=', '00123')]])

    def test_each_expense_line_has_its_own_balanced_posted_move(self):
        class Base:
            def ensure_one(self):
                pass
        cls = load_methods('account_expense_transaction/models/account_expense_transaction.py', 'ExpenseTransaction',
                           {'create_expense_journal_entry'}, Base)
        expense = cls()
        captured, posted = [], []
        class Moves:
            def __init__(self, ids=()):
                self.ids = list(ids)
            def create(self, values):
                captured.append(values)
                return Moves([len(captured)])
            def __or__(self, other):
                return Moves(self.ids + other.ids)
            def action_post(self):
                posted.extend(self.ids)
        def detail(amount):
            return NS(description='Expense', price_total=amount, prepaid_expense_account_id=NS(id=1),
                      analytic_distribution={'10,20': 100}, operating_unit_id=NS(id=8),
                      partner_id=NS(id=7), customer_code='C1', customer_account_domain='[]',
                      customer_account=NS(id=6), contract_type='Type')
        expense.env = {'account.move': Moves()}
        expense.expense_detail_ids = [detail(100), detail(200)]
        expense.company_id, expense.payment_journal_id = NS(id=3), NS(id=4)
        expense.journal_id, expense.bank_account_id = NS(id=4, default_account_id=NS(id=5)), NS(id=5)
        expense.type_jv, expense.date, expense.reference = 'each_line', '2026-10-01', 'EXP1'
        expense.create_expense_journal_entry()
        self.assertEqual(posted, [1, 2])
        self.assertEqual([len(v['line_ids']) for v in captured], [2, 2])
        for values in captured:
            self.assertEqual(values['company_id'], 3)
            self.assertEqual(sum(v.get('debit', 0) - v.get('credit', 0) for _, _, v in values['line_ids']), 0)

    def test_user_defaults_do_not_shadow_fields_module(self):
        class Base:
            def default_get(self, fields_list):
                return {}
        cls = load_methods('operating_unit/models/res_users.py', 'ResUsers', {'default_get'}, Base,
                           fields=NS(Command=NS(set=lambda ids: (6, 0, ids))))
        user = cls()
        defaults = NS(default_operating_unit_id=NS(id=8), operating_unit_ids=NS(ids=[8, 9]))
        config = NS(sudo=lambda: config, get_param=lambda *a: 'True')
        class Env(dict):
            def ref(self, name, raise_if_not_found=True):
                return defaults
        user.env = Env({'ir.config_parameter': config})
        self.assertEqual(user.default_get(['operating_unit_ids'])['operating_unit_ids'], [(6, 0, [8, 9])])

    def test_view_optional_pre_locate(self):
        cls = load_methods('base_view_inheritance_extension/models/ir_ui_view.py', 'IrUiView', {'apply_inheritance_specs'})
        view = cls()
        calls = []
        view._iter_inheritance_specs = lambda tree: iter([(tree, lambda source, spec: source)])
        source, spec = object(), object()
        self.assertIs(view.apply_inheritance_specs(source, spec, pre_locate=None), source)
        view.apply_inheritance_specs(source, spec, pre_locate=calls.append)
        self.assertEqual(calls, [spec])

    def test_petty_cash_multi_create(self):
        class Base:
            def create(self, vals_list):
                return vals_list
        cls = load_methods('petty_cash_purchases/models/petty_cash_purchases.py', 'PettyCashPurchase', {'create'}, Base)
        purchase = cls()
        names = iter(['PCP1', 'PCP2'])
        purchase.env = {'ir.sequence': NS(next_by_code=lambda code: next(names))}
        values = [{}, {}, {'name': 'Existing'}]
        self.assertEqual(purchase.create(values), [{'name': 'PCP1'}, {'name': 'PCP2'}, {'name': 'Existing'}])

    def test_mapping_batch_delete_guard(self):
        class Base:
            def unlink(self):
                return True
            def __iter__(self):
                return iter(self.rows)
        cls = load_methods('employees_ext/models/pc_ou_mapping.py', 'PcOuMapping', {'unlink'}, Base,
                           ValidationError=ValueError, _=lambda value: value)
        mappings = cls()
        mappings.rows = [NS(project_code='100'), NS(project_code='DEFAULT')]
        with self.assertRaises(ValueError):
            mappings.unlink()
        mappings.rows = [NS(project_code='100'), NS(project_code='150')]
        self.assertTrue(mappings.unlink())

    def test_api_flat_body_and_jsonrpc_params(self):
        path = ROOT / 'contacts_ext/controllers/api_utils.py'
        spec = importlib.util.spec_from_file_location('api_utils', path)
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        data = {'employee_code': 'E01', 'customer_account': 'C01'}
        self.assertEqual(module.parse_api_payload(json.dumps(data)), data)
        self.assertEqual(module.parse_api_payload(json.dumps({'jsonrpc': '2.0', 'params': data})), data)
        self.assertEqual(module.parse_api_payload(json.dumps([data])), [data])
        with self.assertRaises(ValueError):
            module.parse_api_payload('{"jsonrpc":"2.0","params":[]}')

    def test_excel_composite_analytic_keys(self):
        for module in ('custom_journal_excel_export', 'journal_voucher_print_export'):
            tree = ast.parse((ROOT / module / 'models/account_move.py').read_bytes())
            expression = next(n.value for n in ast.walk(tree) if isinstance(n, ast.Assign)
                              and any(isinstance(t, ast.Name) and t.id == 'analytic_ids' for t in n.targets))
            value = eval(compile(ast.Expression(expression), module, 'eval'),
                         {'line': NS(analytic_distribution={'10,20': 50, '10,30': 50})})
            self.assertEqual(value, [10, 20, 30])

    def test_statement_pdf_empty_guard_and_action(self):
        class Base:
            def ensure_one(self):
                pass
        cls = load_methods('account_ext/wizard/full_employee_statement_wizard.py', 'FullEmployeeStatementWizard',
                           {'action_print_pdf'}, Base, UserError=ValueError, _=lambda value: value)
        wizard = cls()
        wizard.line_ids = []
        with self.assertRaises(ValueError):
            wizard.action_print_pdf()
        wizard.line_ids = [object()]
        wizard.env = NS(ref=lambda ref: NS(report_action=lambda obj: {'type': 'ir.actions.report', 'ref': ref}))
        self.assertEqual(wizard.action_print_pdf()['ref'], 'account_ext.action_full_employee_statement_pdf')

    def test_prepaid_posting_copies_identifiers_and_dimensions(self):
        class Base:
            def action_post(self):
                return 'posted'
            def __iter__(self):
                return iter(self.rows)
        cls = load_methods('account_expense_transaction/models/account_move.py', 'AccountMove', {'action_post'}, Base,
                           fields=NS(Date=NS(today=lambda: datetime.date(2026, 10, 1))))
        captured, details = [], []
        expense_type = NS(id=3, prepaid_expense_account_id=NS(id=4), expense_account_id=NS(id=5))
        line = NS(name='Prepaid', expense_type_id=expense_type, analytic_distribution={'10,20': 100},
                  start_date='2026-01-01', end_date='2026-12-31', quantity=1, price_unit=100,
                  employee_id=NS(id=9), operating_unit_id=NS(id=8), partner_id=NS(id=7),
                  customer_code='C1', customer_account=NS(id=6), contract_type='Type')
        move = NS(id=42, company_id=NS(id=2), journal_id=NS(id=1), partner_id=NS(id=7), name='BILL/42')
        transaction = NS(expense_detail_ids=NS(create=details.append), id=55)
        def create(values):
            captured.append(values)
            return transaction
        record = cls()
        record.rows = [move]
        record.env = {'account.move.line': NS(search=lambda domain: [line]),
                      'account.expense.transaction': NS(create=create)}
        self.assertEqual(record.action_post(), 'posted')
        self.assertEqual(captured[0]['move_id'], 42)
        self.assertEqual(captured[0]['company_id'], 2)
        self.assertEqual(captured[0]['expense_nature'], 'prepaid')
        self.assertEqual(details[0]['analytic_distribution'], {'10,20': 100})
        self.assertEqual(details[0]['operating_unit_id'], 8)
        self.assertEqual(details[0]['employee_id'], 9)


if __name__ == '__main__':
    unittest.main()
